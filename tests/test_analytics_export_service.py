from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import AsyncMock, patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from bot.repositories.analytics_report_repository import EventRow
from bot.services import analytics_export_service
from bot.services.analytics_export_service import (
    AnalyticsExport,
    AnalyticsExportService,
    export_filename,
    render_xlsx,
)
from bot.services.analytics_report_service import AnalyticsReport, DailyStats

TZ = "Asia/Dushanbe"


def make_report() -> AnalyticsReport:
    return AnalyticsReport(
        days=2,
        since=datetime(2026, 9, 26, tzinfo=ZoneInfo(TZ)),
        timezone=TZ,
        total_users=120,
        new_users=5,
        bot_started=7,
        checks=3,
        success=2,
        failed=1,
        failed_by_reason={"not_found": 1},
        daily=[
            DailyStats(date(2026, 9, 26), new_users=2, checks=1, success=1, failed=0),
            DailyStats(date(2026, 9, 27), new_users=3, checks=2, success=1, failed=1),
        ],
    )


def make_event(serial: str = "ABC123") -> EventRow:
    # 20:30 UTC is 01:30 next day in Dushanbe (UTC+5).
    return EventRow(
        created_at=datetime(2026, 9, 26, 20, 30, tzinfo=timezone.utc),
        telegram_id=555,
        username="buyer",
        event_type="serial_check_failed",
        serial_number=serial,
        reason="not_found",
        batch_id=None,
    )


def load(export: AnalyticsExport):
    return load_workbook(BytesIO(render_xlsx(export)))


def test_export_filename_covers_period():
    assert export_filename(make_report()) == "stats-20260926-20260927.xlsx"


def test_render_xlsx_has_three_sheets_with_headers():
    workbook = load(AnalyticsExport(make_report(), [make_event()]))

    assert workbook.sheetnames == ["Сводка", "По дням", "События"]
    assert workbook["Сводка"]["A1"].value == "Показатель"
    assert workbook["Сводка"]["A1"].font.bold
    assert workbook["События"].freeze_panes == "A2"


def test_render_xlsx_summary_and_daily_values():
    workbook = load(AnalyticsExport(make_report()))

    summary = dict(workbook["Сводка"].iter_rows(min_row=2, values_only=True))
    assert summary["Пользователей всего"] == 120
    assert summary["Проверок"] == 3
    assert summary["  — не найден"] == 1
    assert "Внимание" not in summary

    daily = list(workbook["По дням"].iter_rows(min_row=2, values_only=True))
    assert daily == [
        (datetime(2026, 9, 26), 2, 1, 1, 0),
        (datetime(2026, 9, 27), 3, 2, 1, 1),
    ]


def test_render_xlsx_events_use_local_time():
    workbook = load(AnalyticsExport(make_report(), [make_event()]))

    rows = list(workbook["События"].iter_rows(min_row=2, values_only=True))
    assert rows == [
        (
            datetime(2026, 9, 27, 1, 30),
            555,
            "buyer",
            "serial_check_failed",
            "ABC123",
            "not_found",
            None,
        )
    ]


def test_render_xlsx_marks_truncation():
    workbook = load(AnalyticsExport(make_report(), truncated=True))

    summary = dict(workbook["Сводка"].iter_rows(min_row=2, values_only=True))
    assert "Внимание" in summary


async def build_export(events: list[EventRow], limit: int) -> AnalyticsExport:
    with (
        patch.object(analytics_export_service, "MAX_EXPORT_EVENTS", limit),
        patch.object(analytics_export_service, "AnalyticsReportService") as report_cls,
        patch.object(analytics_export_service, "AnalyticsReportRepository") as repo_cls,
    ):
        report_cls.return_value.build = AsyncMock(return_value=make_report())
        repo_cls.return_value.list_events = AsyncMock(return_value=events)
        export = await AnalyticsExportService(session=object(), tz=TZ).build(2)
        repo_cls.return_value.list_events.assert_awaited_once_with(make_report().since, limit + 1)
        return export


async def test_build_not_truncated_within_limit():
    export = await build_export([make_event("A"), make_event("B")], limit=2)

    assert len(export.events) == 2
    assert export.truncated is False


async def test_build_truncates_extra_row():
    export = await build_export([make_event("A"), make_event("B"), make_event("C")], limit=2)

    assert [event.serial_number for event in export.events] == ["A", "B"]
    assert export.truncated is True
