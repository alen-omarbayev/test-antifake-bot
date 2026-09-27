import html
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from itertools import chain
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from bot.repositories.serial_number_repository import SerialNumberRepository
from bot.services.normalization import normalize_serial

MAX_SERIAL_LENGTH = 64
CHUNK_SIZE = 1000
INVALID_EXAMPLES_LIMIT = 5

_TEXT_SEPARATORS = re.compile(r"[\s,;]+")


@dataclass
class ParsedSerials:
    rows: list[dict] = field(default_factory=list)
    duplicates: int = 0
    invalid: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ImportResult:
    inserted: int
    existing: int


def _collect(pairs: Iterable[tuple[str, str | None]]) -> ParsedSerials:
    parsed = ParsedSerials()
    seen: set[str] = set()
    for raw_serial, batch_id in pairs:
        serial = normalize_serial(raw_serial)
        if not serial:
            continue
        if len(serial) > MAX_SERIAL_LENGTH:
            parsed.invalid.append(serial)
            continue
        if serial in seen:
            parsed.duplicates += 1
            continue
        seen.add(serial)
        parsed.rows.append({"serial_number": serial, "batch_id": batch_id})
    return parsed


def parse_text(text: str, batch_id: str | None) -> ParsedSerials:
    return _collect((token, batch_id) for token in _TEXT_SEPARATORS.split(text))


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    # Excel stores numeric serials as floats: 123456 comes back as 123456.0.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_xlsx(data: bytes, batch_id: str | None) -> ParsedSerials:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            return ParsedSerials()

        header = [_cell_to_str(cell).lower() for cell in first]
        if "serial_number" in header:
            serial_idx = header.index("serial_number")
            batch_idx = header.index("batch_id") if "batch_id" in header else None
            data_rows = rows
        else:
            serial_idx, batch_idx = 0, None
            data_rows = chain([first], rows)

        def pairs():
            for row in data_rows:
                if not row or serial_idx >= len(row):
                    continue
                file_batch = (
                    _cell_to_str(row[batch_idx])
                    if batch_idx is not None and batch_idx < len(row)
                    else ""
                )
                yield _cell_to_str(row[serial_idx]), file_batch or batch_id

        return _collect(pairs())
    finally:
        workbook.close()


def make_batch_id(tz: str, now: datetime | None = None) -> str:
    local = (now or datetime.now(ZoneInfo(tz))).astimezone(ZoneInfo(tz))
    return f"tg-{local:%Y%m%d-%H%M}"


class SerialImportService:
    def __init__(self, session: AsyncSession) -> None:
        self._repo = SerialNumberRepository(session)

    async def import_rows(self, rows: list[dict]) -> ImportResult:
        inserted = 0
        for start in range(0, len(rows), CHUNK_SIZE):
            inserted += await self._repo.bulk_upsert(rows[start : start + CHUNK_SIZE])
        return ImportResult(inserted=inserted, existing=len(rows) - inserted)


def format_import_result(parsed: ParsedSerials, result: ImportResult) -> str:
    text = (
        f"✅ Добавлено новых: {result.inserted}, уже были в базе: {result.existing}, "
        f"повторов: {parsed.duplicates}, некорректных: {len(parsed.invalid)}"
    )
    if parsed.invalid:
        # Examples echo user input and the bot sends HTML, so they are escaped.
        examples = ", ".join(
            html.escape(serial[:20]) + "…" for serial in parsed.invalid[:INVALID_EXAMPLES_LIMIT]
        )
        text +=f"\nНекорректные (длиннее {MAX_SERIAL_LENGTH} символов): {examples}"
    return text
