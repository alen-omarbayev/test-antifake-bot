from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from bot.repositories.analytics_report_repository import AnalyticsReportRepository, EventRow
from bot.services.analytics_report_service import (
    FAILED_REASON_LABELS,
    AnalyticsReport,
    AnalyticsReportService,
)

# Keeps the workbook comfortably below Telegram's 50 MB upload limit.
MAX_EXPORT_EVENTS = 100_000

_BOLD = Font(bold=True)


@dataclass
class AnalyticsExport:
    report: AnalyticsReport
    events: list[EventRow] = field(default_factory=list)
    truncated: bool = False


class AnalyticsExportService:
    def __init__(self, session: AsyncSession, tz: str) -> None:
        self._report_service = AnalyticsReportService(session, tz)
        self._repo = AnalyticsReportRepository(session)

    async def build(self, days: int, now: datetime | None = None) -> AnalyticsExport:
        report = await self._report_service.build(days, now)
        # One extra row tells us the period has more events than we export.
        events = await self._repo.list_events(report.since, MAX_EXPORT_EVENTS + 1)
        return AnalyticsExport(
            report=report,
            events=events[:MAX_EXPORT_EVENTS],
            truncated=len(events) > MAX_EXPORT_EVENTS,
        )


def period_end(report: AnalyticsReport) -> date:
    return report.daily[-1].day if report.daily else report.since.date()


def export_filename(report: AnalyticsReport) -> str:
    return f"stats-{report.since:%Y%m%d}-{period_end(report):%Y%m%d}.xlsx"


def _sheet(workbook: Workbook, title: str, headers: list[str], widths: list[int]):
    sheet = workbook.create_sheet(title)
    # Write-only sheets need layout set before the first row is appended.
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = _BOLD
        header_cells.append(cell)
    sheet.append(header_cells)
    return sheet


def _summary_rows(export: AnalyticsExport) -> list[tuple[str, object]]:
    report = export.report
    rows: list[tuple[str, object]] = [
        ("Период с", report.since.date()),
        ("Период по", period_end(report)),
        ("Дней", report.days),
        ("Часовой пояс", report.timezone),
        ("Пользователей всего", report.total_users),
        ("Новых пользователей", report.new_users),
        ("Запусков /start", report.bot_started),
        ("Проверок", report.checks),
        ("Успешных", report.success),
        ("Неуспешных", report.failed),
    ]
    rows.extend(
        (f"  — {FAILED_REASON_LABELS.get(reason, reason)}", count)
        for reason, count in sorted(report.failed_by_reason.items())
    )
    if export.truncated:
        rows.append(("Внимание", f"на листе «События» только первые {MAX_EXPORT_EVENTS}"))
    return rows


def render_xlsx(export: AnalyticsExport) -> bytes:
    report = export.report
    zone = ZoneInfo(report.timezone)
    workbook = Workbook(write_only=True)

    summary = _sheet(workbook, "Сводка", ["Показатель", "Значение"], [24, 20])
    for row in _summary_rows(export):
        summary.append(row)

    daily = _sheet(
        workbook,
        "По дням",
        ["Дата", "Новых пользователей", "Проверок", "Успешных", "Неуспешных"],
        [12, 20, 12, 12, 12],
    )
    for stats in report.daily:
        daily.append([stats.day, stats.new_users, stats.checks, stats.success, stats.failed])

    events = _sheet(
        workbook,
        "События",
        ["Время", "Telegram ID", "Username", "Событие", "Серийный номер", "Причина", "Партия"],
        [20, 14, 20, 22, 24, 12, 20],
    )
    for event in export.events:
        # Excel has no timezone support, so store local wall-clock time.
        local_time = event.created_at.astimezone(zone).replace(tzinfo=None)
        events.append(
            [
                local_time,
                event.telegram_id,
                event.username,
                event.event_type,
                event.serial_number,
                event.reason,
                event.batch_id,
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
